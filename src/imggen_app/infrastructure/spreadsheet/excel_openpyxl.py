import openpyxl
from pathlib import Path
from typing import Iterator, Optional
import logging
from imggen_app.domain.models import ProductRow
from imggen_app.domain.exceptions import SpreadsheetError

logger = logging.getLogger(__name__)

HEADERS = ["Product Description", "Image Base Name"]

class ExcelHandler:
    """
    Static utility class to handle Excel (.xlsx) file operations using openpyxl.
    """
    @staticmethod
    def create_template(path: Path) -> None:
        """
        Creates a new Excel file with the standard headers at the specified path.
        Used when the user clicks 'Download Template'.
        """
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Products"
            # Append headers defined in constant
            ws.append(HEADERS)
            wb.save(path)
            logger.info(f"Template created at {path}")
        except Exception as e:
            raise SpreadsheetError(f"Failed to create template: {e}")

    @staticmethod
    def read_rows(path: Path) -> Iterator[ProductRow]:
        """
        Reads data rows from an Excel file, skipping empty lines and validating structure.
        Yields ProductRow domain models for processing.
        """
        try:
            # Load workbook in data_only mode to ignore formulas
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            
            # 1. Validate headers (smart check: case insensitive)
            first_row = [cell.value for cell in ws[1][:2]]
            if not all(h.lower() == r.lower() for h, r in zip(HEADERS, first_row) if r):
                raise SpreadsheetError(f"Invalid headers. Expected: {HEADERS}")
            
            # 2. Iterate through rows starting from line 2
            for i, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
                description, base_name = row
                
                # Skip rows that are completely empty
                if not description and not base_name:
                    continue
                
                # Basic validation: both fields are mandatory for a product
                if not description or not base_name:
                    logger.warning(f"Skipping row {i}: missing data. Desc: '{description}', Name: '{base_name}'")
                    continue
                
                # Clean and yield the row
                yield ProductRow(
                    row_index=i,
                    description=str(description).strip(),
                    base_name=str(base_name).strip()
                )
        except SpreadsheetError:
            # Pass-through for domain-specific errors
            raise
        except Exception as e:
            # Wrap low-level openpyxl or file system errors
            raise SpreadsheetError(f"Failed to read spreadsheet: {e}")
